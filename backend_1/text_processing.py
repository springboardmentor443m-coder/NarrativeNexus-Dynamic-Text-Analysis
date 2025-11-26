from transformers import pipeline
from bs4 import BeautifulSoup
import re
import pdfplumber
import docx
from pptx import Presentation

# Load summarizer and sentiment once
print("🔄 Loading summarizer and sentiment models...")
summarizer = pipeline("summarization", model="facebook/bart-large-cnn")
sentiment_analyzer = pipeline("sentiment-analysis")
print("✅ Summarizer and sentiment ready")


def extract_text_from_file(filepath: str) -> str:
    """Extract text for txt, html, pdf, docx, pptx. Returns empty string if extraction fails."""
    ext = filepath.lower().split(".")[-1]

    if ext in ("txt", "text"):
        try:
            with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        except Exception:
            return ""

    if ext in ("html", "htm"):
        try:
            with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        except Exception:
            return ""

    if ext == "pdf":
        try:
            text = ""
            with pdfplumber.open(filepath) as pdf:
                for page in pdf.pages:
                    text += (page.extract_text() or "") + "\n"
            return text
        except Exception:
            return ""

    if ext == "docx":
        try:
            doc = docx.Document(filepath)
            return "\n".join([p.text for p in doc.paragraphs])
        except Exception:
            return ""

    if ext == "pptx":
        try:
            prs = Presentation(filepath)
            texts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        texts.append(shape.text)
            return "\n".join(texts)
        except Exception:
            return ""

    # Fallback: try to read as text
    try:
        with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    except Exception:
        return ""


def clean_text(raw_text: str) -> str:
    """Remove HTML tags, URLs, odd characters, and normalize whitespace."""
    soup = BeautifulSoup(raw_text, "html.parser")
    text = soup.get_text(separator=" ")

    text = re.sub(r"http\S+|www\S+", "", text)
    # keep basic punctuation + alphanumerics and whitespace
    text = re.sub(r"[^a-zA-Z0-9\.\,\!\?\;\:\'\"\(\)\-\s]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()

    return text


def summarize_text(text: str, max_length: int = 150, min_length: int = 40) -> str:
    """Summarize using BART. Truncate long inputs to 3000 chars to avoid token limits."""
    if not text or not text.strip():
        return "No valid text provided."

    truncated = text[:3000]
    word_count = len(truncated.split())

    if word_count < 100:
        max_length, min_length = 60, 20
    elif word_count < 300:
        max_length, min_length = 120, 40
    else:
        max_length, min_length = 150, 40

    summary = summarizer(truncated, max_length=max_length, min_length=min_length, do_sample=False)[0]["summary_text"]
    return summary.strip()


def analyze_sentiment(text: str) -> dict:
    """Return sentiment label and score (rounded)."""
    if not text or not text.strip():
        return {"label": "NEUTRAL", "score": 0.0}

    result = sentiment_analyzer(text[:512])[0]
    return {"label": result["label"], "score": round(result["score"], 3)}
